# -*- coding: utf-8 -*-
"""Diagnose a .docx: heading structure, styles, tables, images.

Usage: python inspect_docx.py <file.docx>
"""
import sys, os, json, collections
try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
from docx import Document
from docx.shared import Pt


def run_info(r):
    sz = r.font.size.pt if r.font.size else None
    return {
        "text": r.text[:80],
        "bold": bool(r.bold), "italic": bool(r.italic),
        "size": sz, "color": str(r.font.color.rgb) if r.font.color and r.font.color.type is not None else None,
    }


def main(path):
    doc = Document(path)
    styles = collections.Counter(p.style.name for p in doc.paragraphs)
    paras = []
    for i, p in enumerate(doc.paragraphs[:60]):
        if not p.text.strip():
            continue
        pf = p.paragraph_format
        paras.append({
            "i": i, "style": p.style.name, "align": str(p.alignment),
            "indent_first": pf.first_line_indent.pt if pf.first_line_indent else None,
            "line_spacing": str(pf.line_spacing),
            "runs": [run_info(r) for r in p.runs[:4]],
            "text": p.text[:100],
        })
    tables = []
    for ti, t in enumerate(doc.tables[:6]):
        tables.append({"rows": len(t.rows), "cols": len(t.columns),
                       "head": [c.text[:20] for c in t.rows[0].cells][:8] if len(t.rows) else []})
    sect = doc.sections[0]
    out = {
        "file": os.path.basename(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "style_usage": styles.most_common(15),
        "page": {"w_pt": sect.page_width.pt, "h_pt": sect.page_height.pt,
                 "margins": [sect.left_margin.pt, sect.top_margin.pt, sect.right_margin.pt, sect.bottom_margin.pt]},
        "paragraphs": paras,
        "tables": tables,
    }
    print(json.dumps(out, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main(sys.argv[1])
